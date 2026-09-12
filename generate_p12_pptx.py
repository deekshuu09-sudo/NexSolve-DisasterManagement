import os
import sys
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette Constants
    COLOR_NAVY_DARK = RGBColor(15, 23, 42)      # #0F172A
    COLOR_NAVY_LIGHT = RGBColor(30, 41, 59)     # #1E293B
    COLOR_BLUE_ACCENT = RGBColor(30, 64, 175)   # #1E40AF
    COLOR_BLUE_BRIGHT = RGBColor(37, 99, 235)   # #2563EB
    COLOR_GOLD = RGBColor(217, 119, 6)          # #D97706
    COLOR_GOLD_LIGHT = RGBColor(245, 158, 11)   # #F59E0B
    COLOR_GREEN = RGBColor(5, 150, 105)         # #059669
    COLOR_SLATE_BG = RGBColor(248, 250, 252)    # #F8FAFC
    COLOR_CARD_BG = RGBColor(255, 255, 255)     # #FFFFFF
    COLOR_TEXT_MAIN = RGBColor(15, 23, 42)     # #0F172A
    COLOR_TEXT_MUTED = RGBColor(71, 85, 105)   # #475569
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_BORDER_GRAY = RGBColor(226, 232, 240) # #E2E8F0

    # Helper: Add Standard Header
    def add_header(slide, title_text, category_text="SMART INDIA HACKATHON 2026"):
        # Header banner shape
        header_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
        header_box.fill.solid()
        header_box.fill.fore_color.rgb = COLOR_NAVY_DARK
        header_box.line.color.rgb = COLOR_NAVY_DARK

        # Top Accent Line
        accent_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.06), Inches(13.333), Inches(0.04))
        accent_line.fill.solid()
        accent_line.fill.fore_color.rgb = COLOR_GOLD
        accent_line.line.color.rgb = COLOR_GOLD

        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.6)
        tf.margin_top = Inches(0.15)
        
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.name = 'Arial'
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_GOLD_LIGHT

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.name = 'Arial'
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_WHITE

        # Footer
        footer_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.1), Inches(13.333), Inches(0.4))
        footer_box.fill.solid()
        footer_box.fill.fore_color.rgb = COLOR_NAVY_LIGHT
        footer_box.line.color.rgb = COLOR_NAVY_LIGHT

        ftf = footer_box.text_frame
        ftf.margin_left = Inches(0.6)
        ftf.margin_top = Inches(0.08)
        fp = ftf.paragraphs[0]
        fp.text = "NexSolve — Landslide Risk Intelligence & Decision Support | Official SIH 2026 Idea Submission Template"
        fp.font.name = 'Arial'
        fp.font.size = Pt(9)
        fp.font.color.rgb = COLOR_WHITE

    # -------------------------------------------------------------
    # SLIDE 1: TITLE PAGE
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_NAVY_DARK
    bg1.line.color.rgb = COLOR_NAVY_DARK

    # Gold Accent Top Bar
    bar1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
    bar1.fill.solid()
    bar1.fill.fore_color.rgb = COLOR_GOLD
    bar1.line.color.rgb = COLOR_GOLD

    # Title Card Text Box
    tb1 = slide1.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.733), Inches(2.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "SMART INDIA HACKATHON 2026"
    p.font.name = 'Arial'
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD_LIGHT

    p = tf1.add_paragraph()
    p.text = "NEXSOLVE — LANDSLIDE RISK INTELLIGENCE &\nDECISION SUPPORT FOR NORTHEAST INDIA"
    p.font.name = 'Arial'
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE

    p = tf1.add_paragraph()
    p.text = "AI-Assisted Geospatial Risk Intelligence & Operational Decision Support System"
    p.font.name = 'Arial'
    p.font.size = Pt(14)
    p.font.italic = True
    p.font.color.rgb = RGBColor(203, 213, 225) # Slate 300

    # Fields Grid Container (2 Columns of Cards)
    fields = [
        ("Problem Statement ID:", "[INSERT SIH PROBLEM STATEMENT ID]"),
        ("Problem Statement Title:", "[INSERT SIH PROBLEM STATEMENT TITLE]"),
        ("Theme:", "Disaster Management"),
        ("PS Category:", "Software"),
        ("Team ID:", "[INSERT SIH TEAM ID]"),
        ("Team Name:", "NexSolve")
    ]

    for idx, (label, val) in enumerate(fields):
        col = idx % 2
        row = idx // 2
        x = Inches(0.8 + col * 5.9)
        y = Inches(3.0 + row * 1.25)
        w = Inches(5.6)
        h = Inches(1.1)

        card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_NAVY_LIGHT
        card.line.color.rgb = COLOR_BLUE_ACCENT
        card.line.width = Pt(1.5)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.25)
        ctf.margin_top = Inches(0.15)

        cp0 = ctf.paragraphs[0]
        cp0.text = label.upper()
        cp0.font.name = 'Arial'
        cp0.font.size = Pt(10)
        cp0.font.bold = True
        cp0.font.color.rgb = COLOR_GOLD_LIGHT

        cp1 = ctf.add_paragraph()
        cp1.text = val
        cp1.font.name = 'Arial'
        cp1.font.size = Pt(13)
        cp1.font.bold = True
        cp1.font.color.rgb = COLOR_WHITE

    # Bottom Footer Banner Slide 1
    ft1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(6.95), Inches(13.333), Inches(0.55))
    ft1.fill.solid()
    ft1.fill.fore_color.rgb = COLOR_BLUE_ACCENT
    ft1.line.color.rgb = COLOR_BLUE_ACCENT
    ftf1 = ft1.text_frame
    ftf1.margin_left = Inches(0.8)
    ftf1.margin_top = Inches(0.12)
    fp1 = ftf1.paragraphs[0]
    fp1.text = "OFFICIAL SIH 2026 IDEA SUBMISSION FORMAT | DECISION-SUPPORT PROTOTYPE FOR NORTHEAST INDIA"
    fp1.font.name = 'Arial'
    fp1.font.size = Pt(10)
    fp1.font.bold = True
    fp1.font.color.rgb = COLOR_WHITE


    # -------------------------------------------------------------
    # SLIDE 2: IDEA TITLE / PROPOSED SOLUTION
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    bg2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg2.fill.solid()
    bg2.fill.fore_color.rgb = COLOR_SLATE_BG
    bg2.line.color.rgb = COLOR_SLATE_BG
    add_header(slide2, "IDEA TITLE / PROPOSED SOLUTION")

    # Idea Title Banner
    ib2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.25), Inches(12.133), Inches(0.65))
    ib2.fill.solid()
    ib2.fill.fore_color.rgb = COLOR_BLUE_ACCENT
    ib2.line.color.rgb = COLOR_BLUE_ACCENT
    itf2 = ib2.text_frame
    itf2.margin_left = Inches(0.3)
    itf2.margin_top = Inches(0.12)
    ip2 = itf2.paragraphs[0]
    ip2.text = "IDEA TITLE: NEXSOLVE — LANDSLIDE RISK INTELLIGENCE & DECISION SUPPORT FOR NORTHEAST INDIA"
    ip2.font.name = 'Arial'
    ip2.font.size = Pt(13)
    ip2.font.bold = True
    ip2.font.color.rgb = COLOR_WHITE

    # 4-Stage Pipeline Banner
    pipeline_steps = [
        "1. MULTI-SOURCE DATA\n(Weather, GeoJSON, DEM, GSI)",
        "2. AI RISK INTELLIGENCE\n(7-Feature RF Model & Open-Meteo)",
        "3. DECISION ENGINE\n(Decoupled Exposure & Watchlist)",
        "4. HUMAN VERIFICATION\n(Situation Room Dashboard)"
    ]
    for idx, pstep in enumerate(pipeline_steps):
        px = Inches(0.6 + idx * 3.08)
        pw = Inches(2.9)
        pshape = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px, Inches(2.0), pw, Inches(0.75))
        pshape.fill.solid()
        pshape.fill.fore_color.rgb = COLOR_NAVY_DARK
        pshape.line.color.rgb = COLOR_GOLD
        pshape.line.width = Pt(1.5)
        ptf = pshape.text_frame
        ptf.margin_left = Inches(0.1)
        ptf.margin_top = Inches(0.1)
        pp = ptf.paragraphs[0]
        pp.text = pstep
        pp.alignment = PP_ALIGN.CENTER
        pp.font.name = 'Arial'
        pp.font.size = Pt(10)
        pp.font.bold = True
        pp.font.color.rgb = COLOR_WHITE

    # Problem vs Solution Left Box
    pb_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(2.9), Inches(5.9), Inches(4.0))
    pb_box.fill.solid()
    pb_box.fill.fore_color.rgb = COLOR_CARD_BG
    pb_box.line.color.rgb = COLOR_BORDER_GRAY
    pb_box.line.width = Pt(1.5)

    ptf2 = pb_box.text_frame
    ptf2.word_wrap = True
    ptf2.margin_left = Inches(0.25)
    ptf2.margin_top = Inches(0.2)

    p = ptf2.paragraphs[0]
    p.text = "THE DISASTER MANAGEMENT CHALLENGE"
    p.font.name = 'Arial'
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD

    bullets_prob = [
        "Fragmented & Siloed Data: Rainfall, historical landslide, and socio-economic exposure datasets exist in isolated formats.",
        "Monsoon Vulnerability: Heavy 1d/3d/7d antecedent rainfall triggers sudden slope failures along arterial highways (NH-306, NH-2, NH-10, NH-6).",
        "Lack of Spatial Context: Generic regional advisories lack Survey of India district-level boundary precision.",
        "Unclear Data Transparency: Operational tools often mask missing exposure data with synthetic zero values."
    ]
    for b in bullets_prob:
        p = ptf2.add_paragraph()
        p.text = "• " + b
        p.font.name = 'Arial'
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_TEXT_MAIN

    # Proposed Solution & Innovation Right Box
    sol_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.9), Inches(5.9), Inches(4.0))
    sol_box.fill.solid()
    sol_box.fill.fore_color.rgb = COLOR_CARD_BG
    sol_box.line.color.rgb = COLOR_BLUE_ACCENT
    sol_box.line.width = Pt(1.5)

    stf2 = sol_box.text_frame
    stf2.word_wrap = True
    stf2.margin_left = Inches(0.25)
    stf2.margin_top = Inches(0.2)

    p = stf2.paragraphs[0]
    p.text = "NEXSOLVE PROPOSED SOLUTION & INNOVATION"
    p.font.name = 'Arial'
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_BRIGHT

    bullets_sol = [
        "Unified Geospatial Risk Intelligence: Combines weather-aware ML with official 131 Survey of India administrative boundaries.",
        "Decoupled Architecture: Hazard probability (ML) is computationally separate from socio-economic exposure and vulnerability.",
        "Data Quality Transparency: Explicit INSUFFICIENT DATA reporting for 115 districts rather than synthetic zero-fill.",
        "Explainable Decision Engine: Local feature attribution, antecedent rainfall drivers, and confidence metrics.",
        "Human-in-the-Loop Protocol: Designed specifically for authority decision-support — NOT automated public alarms/evacuation."
    ]
    for b in bullets_sol:
        p = stf2.add_paragraph()
        p.text = "• " + b
        p.font.name = 'Arial'
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_TEXT_MAIN


    # -------------------------------------------------------------
    # SLIDE 3: TECHNICAL APPROACH
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    bg3 = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg3.fill.solid()
    bg3.fill.fore_color.rgb = COLOR_SLATE_BG
    bg3.line.color.rgb = COLOR_SLATE_BG
    add_header(slide3, "TECHNICAL APPROACH & SYSTEM ARCHITECTURE")

    # Workflow Diagram Row (5 Process Nodes)
    nodes = [
        ("DATA INGESTION", "Open-Meteo Weather API\nGSI 10.49k History\nSoI 131 GeoJSON\nUSGS/AWS 30m DEM"),
        ("FEATURE PIPELINE", "1d, 3d, 7d Antecedent Rain\nSpatial Coordinates\nCyclic Month Encoding\nDEM Spatial Context"),
        ("ML HAZARD MODEL", "Random Forest\n(candidate_a_rf_v1)\n7 Production Features\nROC-AUC: 0.9206"),
        ("DECISION ENGINE", "Decoupled Exposure Engine\nMulti-Criterion Watchlist\nFeature Attribution\nConfidence Scoring"),
        ("OPERATIONAL UI", "React + Vite + TS Dashboard\nInteractive Leaflet GIS\nData Quality Audit\nHuman Verification")
    ]

    for idx, (title, desc) in enumerate(nodes):
        nx = Inches(0.6 + idx * 2.45)
        nw = Inches(2.3)
        nh = Inches(2.2)
        nbox = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, nx, Inches(1.3), nw, nh)
        nbox.fill.solid()
        nbox.fill.fore_color.rgb = COLOR_NAVY_DARK if idx == 2 else COLOR_CARD_BG
        nbox.line.color.rgb = COLOR_GOLD if idx == 2 else COLOR_BLUE_ACCENT
        nbox.line.width = Pt(1.5)

        ntf = nbox.text_frame
        ntf.word_wrap = True
        ntf.margin_left = Inches(0.15)
        ntf.margin_top = Inches(0.15)

        p = ntf.paragraphs[0]
        p.text = f"0{idx+1}. {title}"
        p.font.name = 'Arial'
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_GOLD_LIGHT if idx == 2 else COLOR_BLUE_BRIGHT

        p = ntf.add_paragraph()
        p.text = desc
        p.font.name = 'Arial'
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_WHITE if idx == 2 else COLOR_TEXT_MAIN

    # Bottom Left: 7 Production ML Features Card
    feat_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(3.7), Inches(5.9), Inches(3.2))
    feat_box.fill.solid()
    feat_box.fill.fore_color.rgb = COLOR_CARD_BG
    feat_box.line.color.rgb = COLOR_BORDER_GRAY
    feat_box.line.width = Pt(1.5)

    ftf3 = feat_box.text_frame
    ftf3.word_wrap = True
    ftf3.margin_left = Inches(0.25)
    ftf3.margin_top = Inches(0.2)

    p = ftf3.paragraphs[0]
    p.text = "PRODUCTION ML MODEL CONTRACT (7 FEATURES)"
    p.font.name = 'Arial'
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_DARK

    f_bullets = [
        "Spatial Coordinates: latitude, longitude",
        "Weather Feeds: rainfall_1d, rainfall_3d, rainfall_7d (antecedent accumulation)",
        "Cyclic Temporal Encoding: month_sin, month_cos",
        "Environmental Context (Decoupled): USGS/AWS NASADEM 30m terrain elevation, slope, aspect statistics provide environmental decision context and are computationally separate from the 7-feature RF model."
    ]
    for b in f_bullets:
        p = ftf3.add_paragraph()
        p.text = "• " + b
        p.font.name = 'Arial'
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_TEXT_MAIN

    # Bottom Right: Evaluation Metrics & Tech Stack
    right_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(3.7), Inches(5.9), Inches(3.2))
    right_box.fill.solid()
    right_box.fill.fore_color.rgb = COLOR_CARD_BG
    right_box.line.color.rgb = COLOR_BLUE_ACCENT
    right_box.line.width = Pt(1.5)

    rtf3 = right_box.text_frame
    rtf3.word_wrap = True
    rtf3.margin_left = Inches(0.25)
    rtf3.margin_top = Inches(0.2)

    p = rtf3.paragraphs[0]
    p.text = "VERIFIED EVALUATION METRICS & TECH STACK"
    p.font.name = 'Arial'
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_BRIGHT

    p = rtf3.add_paragraph()
    p.text = "2024–2025 Locked Temporal Holdout Evaluation:"
    p.font.name = 'Arial'
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD

    metrics_str = "• ROC-AUC: 0.9206  |  PR-AUC: 0.9328\n• Precision: 0.8312  |  Recall: 0.8204  |  F1: 0.8258\n*(Evaluated on strict temporal holdout; spatial data leakage discarded)*"
    p = rtf3.add_paragraph()
    p.text = metrics_str
    p.font.name = 'Arial'
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_TEXT_MAIN

    p = rtf3.add_paragraph()
    p.text = "Technology Stack:"
    p.font.name = 'Arial'
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY_DARK

    t_str = "• Frontend: React 18, Vite, TypeScript, Leaflet GIS, Tailwind CSS\n• Backend: Python 3, FastAPI microservices, GeoPandas, Shapely\n• Data Sources: Open-Meteo Weather API, Survey of India, USGS/AWS DEM"
    p = rtf3.add_paragraph()
    p.text = t_str
    p.font.name = 'Arial'
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_TEXT_MAIN


    # -------------------------------------------------------------
    # SLIDE 4: FEASIBILITY AND VIABILITY
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    bg4 = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg4.fill.solid()
    bg4.fill.fore_color.rgb = COLOR_SLATE_BG
    bg4.line.color.rgb = COLOR_SLATE_BG
    add_header(slide4, "FEASIBILITY, RISKS & TRANSPARENT COVERAGE")

    # 3 Column Layout (Feasible Now / Challenges & Risks / Mitigation Strategies)
    cols_data = [
        ("FEASIBLE NOW (Working Prototype)", COLOR_GREEN, [
            "Functional React + FastAPI microservice architecture",
            "Official 131-district Survey of India boundary integration",
            "Live Open-Meteo weather integration & 1d/3d/7d rain tracking",
            "ML hazard risk inference (7-feature RF engine)",
            "USGS/AWS NASADEM 30m terrain feature processing",
            "Decoupled exposure intelligence & situation room dashboard"
        ]),
        ("CHALLENGES & RISKS", COLOR_GOLD, [
            "Incomplete district exposure data across NE states",
            "Historical landslide inventory spatial/temporal reporting gaps",
            "Live weather station density and quality variations",
            "No official operational NDMA/SDMA warning integration yet",
            "Public cloud deployment not yet verified (container-ready)"
        ]),
        ("MITIGATION STRATEGIES", COLOR_BLUE_BRIGHT, [
            "Transparent INSUFFICIENT DATA reporting (no synthetic zero-fill)",
            "Data provenance tracking & explicit freshness indicators",
            "Temporal holdout validation (preventing data leakage)",
            "Confidence & data completeness separated from probability",
            "Human-in-the-loop verification before decision action"
        ])
    ]

    for idx, (ctitle, ccolor, cbullets) in enumerate(cols_data):
        cx = Inches(0.6 + idx * 4.08)
        cw = Inches(3.9)
        ch = Inches(4.3)
        cbox = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(1.3), cw, ch)
        cbox.fill.solid()
        cbox.fill.fore_color.rgb = COLOR_CARD_BG
        cbox.line.color.rgb = ccolor
        cbox.line.width = Pt(1.5)

        ctf4 = cbox.text_frame
        ctf4.word_wrap = True
        ctf4.margin_left = Inches(0.2)
        ctf4.margin_top = Inches(0.2)

        p = ctf4.paragraphs[0]
        p.text = ctitle
        p.font.name = 'Arial'
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ccolor

        for b in cbullets:
            p = ctf4.add_paragraph()
            p.text = "• " + b
            p.font.name = 'Arial'
            p.font.size = Pt(9.5)
            p.font.color.rgb = COLOR_TEXT_MAIN

    # Bottom Transparency & Deployment Banner (2 Sub Cards)
    # Card 1: Vulnerability Coverage Transparency
    cov_card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(5.75), Inches(7.0), Inches(1.2))
    cov_card.fill.solid()
    cov_card.fill.fore_color.rgb = COLOR_NAVY_DARK
    cov_card.line.color.rgb = COLOR_GOLD
    cov_card.line.width = Pt(1.5)
    ctf_cov = cov_card.text_frame
    ctf_cov.margin_left = Inches(0.2)
    ctf_cov.margin_top = Inches(0.12)
    p = ctf_cov.paragraphs[0]
    p.text = "GEOGRAPHIC & VULNERABILITY COVERAGE TRANSPARENCY"
    p.font.name = 'Arial'
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD_LIGHT
    p = ctf_cov.add_paragraph()
    p.text = "• 131 Districts = Official Survey of India Geographic Framework\n• 16 Scored Districts (8 COMPUTED full profiles, 8 PARTIAL) | 115 Districts = INSUFFICIENT_DATA (score null)\n• Governance Rule: Missing exposure data is NEVER converted to synthetic zeros."
    p.font.name = 'Arial'
    p.font.size = Pt(9.5)
    p.font.color.rgb = COLOR_WHITE

    # Card 2: Deployment Status Transparency
    dep_card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(5.75), Inches(4.933), Inches(1.2))
    dep_card.fill.solid()
    dep_card.fill.fore_color.rgb = COLOR_NAVY_LIGHT
    dep_card.line.color.rgb = COLOR_BLUE_ACCENT
    dep_card.line.width = Pt(1.5)
    ctf_dep = dep_card.text_frame
    ctf_dep.margin_left = Inches(0.2)
    ctf_dep.margin_top = Inches(0.12)
    p = ctf_dep.paragraphs[0]
    p.text = "DEPLOYMENT READINESS VERDICT"
    p.font.name = 'Arial'
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_BRIGHT
    p = ctf_dep.add_paragraph()
    p.text = "• Deployment Wording: \"Container-ready; public cloud deployment not yet verified.\"\n• System Readiness: 150/150 backend unit tests PASS | Frontend production build PASS."
    p.font.name = 'Arial'
    p.font.size = Pt(9.5)
    p.font.color.rgb = COLOR_WHITE


    # -------------------------------------------------------------
    # SLIDE 5: IMPACT AND BENEFITS
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    bg5 = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg5.fill.solid()
    bg5.fill.fore_color.rgb = COLOR_SLATE_BG
    bg5.line.color.rgb = COLOR_SLATE_BG
    add_header(slide5, "POTENTIAL IMPACT & STAKEHOLDER BENEFITS")

    # Value Proposition Top Banner
    vbanner = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.25), Inches(12.133), Inches(0.65))
    vbanner.fill.solid()
    vbanner.fill.fore_color.rgb = COLOR_BLUE_ACCENT
    vbanner.line.color.rgb = COLOR_GOLD
    vbanner.line.width = Pt(1.5)
    vtf = vbanner.text_frame
    vtf.margin_left = Inches(0.3)
    vtf.margin_top = Inches(0.12)
    vp = vtf.paragraphs[0]
    vp.text = "KEY VALUE PROPOSITION: \"FROM FRAGMENTED DISASTER DATA TO ACTIONABLE RISK INTELLIGENCE.\""
    vp.font.name = 'Arial'
    vp.font.size = Pt(13)
    vp.font.bold = True
    vp.alignment = PP_ALIGN.CENTER
    vp.font.color.rgb = COLOR_WHITE

    # 4 Target Stakeholder Grid Cards
    stakeholders = [
        ("DISASTER MANAGEMENT AUTHORITIES", COLOR_BLUE_BRIGHT, [
            "State EOC & SDMA decision support",
            "Risk-based resource allocation & prioritization",
            "District watchlists for high-antecedent rainfall zones"
        ]),
        ("DISTRICT ADMINISTRATION", COLOR_GOLD, [
            "DDMA preparedness planning",
            "Identification of vulnerable administrative sectors",
            "Contextual critical infrastructure & asset exposure"
        ]),
        ("EMERGENCY RESPONSE TEAMS", COLOR_GREEN, [
            "NDRF / SDRF situational awareness",
            "Arterial highway corridor monitoring (NH-306, NH-2, NH-10)",
            "Pre-positioning equipment near high-hazard slopes"
        ]),
        ("LOCAL COMMUNITIES", COLOR_NAVY_LIGHT, [
            "Future potential access to verified advisories",
            "Transparent confidence and data quality indicators",
            "Enhanced community preparedness awareness"
        ])
    ]

    for idx, (stitle, scolor, sbullets) in enumerate(stakeholders):
        col = idx % 2
        row = idx // 2
        sx = Inches(0.6 + col * 6.18)
        sy = Inches(2.05 + row * 2.05)
        sw = Inches(5.95)
        sh = Inches(1.9)

        sbox = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx, sy, sw, sh)
        sbox.fill.solid()
        sbox.fill.fore_color.rgb = COLOR_CARD_BG
        sbox.line.color.rgb = scolor
        sbox.line.width = Pt(1.5)

        stf5 = sbox.text_frame
        stf5.word_wrap = True
        stf5.margin_left = Inches(0.2)
        stf5.margin_top = Inches(0.15)

        p = stf5.paragraphs[0]
        p.text = stitle
        p.font.name = 'Arial'
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = scolor

        for b in sbullets:
            p = stf5.add_paragraph()
            p.text = "• " + b
            p.font.name = 'Arial'
            p.font.size = Pt(9.5)
            p.font.color.rgb = COLOR_TEXT_MAIN

    # Bottom 3 Impact Dimension Cards
    dimensions = [
        ("SOCIAL IMPACT", "Faster prioritization, enhanced situational awareness, transparent confidence indicators."),
        ("ECONOMIC IMPACT", "Targeted infrastructure protection, optimized response funding, minimized highway disruption."),
        ("ENVIRONMENTAL IMPACT", "Geospatial terrain understanding, slope vulnerability context, targeted mitigation support.")
    ]
    for idx, (dtitle, ddesc) in enumerate(dimensions):
        dx = Inches(0.6 + idx * 4.08)
        dw = Inches(3.9)
        dh = Inches(0.85)
        dbox = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, dx, Inches(6.15), dw, dh)
        dbox.fill.solid()
        dbox.fill.fore_color.rgb = COLOR_NAVY_DARK
        dbox.line.color.rgb = COLOR_BLUE_ACCENT

        dtf = dbox.text_frame
        dtf.word_wrap = True
        dtf.margin_left = Inches(0.15)
        dtf.margin_top = Inches(0.1)

        p = dtf.paragraphs[0]
        p.text = dtitle
        p.font.name = 'Arial'
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_GOLD_LIGHT

        p = dtf.add_paragraph()
        p.text = ddesc
        p.font.name = 'Arial'
        p.font.size = Pt(8.5)
        p.font.color.rgb = COLOR_WHITE


    # -------------------------------------------------------------
    # SLIDE 6: RESEARCH AND REFERENCES
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    bg6 = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg6.fill.solid()
    bg6.fill.fore_color.rgb = COLOR_SLATE_BG
    bg6.line.color.rgb = COLOR_SLATE_BG
    add_header(slide6, "RESEARCH, DATA PROVENANCE & REFERENCES")

    sources = [
        ("Survey of India (SoI)", "Administrative Boundary Database (ABDB LGD Integrated 131 Districts)", "Official Administrative Geometry"),
        ("Geological Survey of India (GSI)", "NLSM & Historical Landslide Inventory (10,492 Records)", "Historical Landslide Baselines"),
        ("USGS / AWS Open Data Elevation Archive", "NASADEM 1 Arc-Second (~30m SRTM DEM) Elevation Grid", "Terrain & Morphometry Features"),
        ("India Meteorological Department (IMD)", "Gridded Rainfall Records (2014–2025)", "Historical Rainfall Calibration"),
        ("Open-Meteo API", "Operational Live Weather & Antecedent Rain Tracking", "Real-Time Weather Feed"),
        ("NITI Aayog", "National Multidimensional Poverty Index (MPI 2023)", "Socio-Economic Vulnerability"),
        ("MoHFW / National Health Authority", "Health Facility Registry (HFR)", "Critical Healthcare Infrastructure"),
        ("Ministry of Education / UDISE+", "Universal District Information System for Education", "School & Emergency Shelter Capacity"),
        ("MoRTH / NHAI GIS", "National Highway Spatial Alignment (NH-306, NH-2, NH-10, NH-6)", "Arterial Transport Corridors"),
        ("BMTPC & Census of India", "Vulnerability Atlas of India & 2011 Primary Census Abstract", "Housing & Demographic Context")
    ]

    for idx, (org, desc, cat) in enumerate(sources):
        col = idx % 2
        row = idx // 5
        rx = Inches(0.6 + col * 6.18)
        ry = Inches(1.3 + row * 1.1)
        rw = Inches(5.95)
        rh = Inches(0.98)

        rbox = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, rx, ry, rw, rh)
        rbox.fill.solid()
        rbox.fill.fore_color.rgb = COLOR_CARD_BG
        rbox.line.color.rgb = COLOR_BORDER_GRAY
        rbox.line.width = Pt(1.5)

        rtf = rbox.text_frame
        rtf.word_wrap = True
        rtf.margin_left = Inches(0.2)
        rtf.margin_top = Inches(0.1)

        p = rtf.paragraphs[0]
        p.text = f"{idx+1}. {org}"
        p.font.name = 'Arial'
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_NAVY_DARK

        p = rtf.add_paragraph()
        p.text = f"Category: {cat}"
        p.font.name = 'Arial'
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = COLOR_GOLD

        p = rtf.add_paragraph()
        p.text = f"Source Details: {desc}"
        p.font.name = 'Arial'
        p.font.size = Pt(8.5)
        p.font.color.rgb = COLOR_TEXT_MUTED

    # Bottom Governance & Scientific Integrity Box
    gbox = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(6.8), Inches(12.133), Inches(0.25))
    gbox.fill.solid()
    gbox.fill.fore_color.rgb = COLOR_NAVY_DARK
    gbox.line.color.rgb = COLOR_NAVY_DARK
    gtf = gbox.text_frame
    gtf.margin_left = Inches(0.2)
    gtf.margin_top = Inches(0.02)
    gp = gtf.paragraphs[0]
    gp.text = "DATA PROVENANCE STATEMENT: All administrative geometries, weather parameters, terrain rasters, and vulnerability indicators are derived exclusively from authoritative government repositories and open scientific archives."
    gp.font.name = 'Arial'
    gp.font.size = Pt(8)
    gp.font.color.rgb = COLOR_WHITE

    # Save PPTX
    out_pptx = "/Users/deekshitharoy/Desktop/NexSolve-DisasterManagement/NexSolve_SIH_Idea_Submission_2026.pptx"
    prs.save(out_pptx)
    print(f"Presentation saved to {out_pptx}")

if __name__ == '__main__':
    build_presentation()
